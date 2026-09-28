import io
import json
import os
import re
import sqlite3
import urllib.request
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, session
from werkzeug.security import generate_password_hash, check_password_hash


BASE_DIR = Path(__file__).resolve().parent
DB_DIR = BASE_DIR / "instance"
DB_PATH = DB_DIR / "learn4all.db"
DB_DIR.mkdir(exist_ok=True)

app = Flask(__name__, template_folder="templates", static_folder="static")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "inclulearn-dev-secret-change-me")

STOP_WORDS = set(
    "және мен бұл үшін туралы немесе бір оның болып арқылы да де әрі қалай деген "
    "the and of to a in is it that for with as on are was this be by an or from at "
    "и в на это как для что из с по".split()
)

LANGUAGE_NAMES = {
    "kk": "қазақ тілі",
    "ru": "русский язык",
    "en": "English",
}


@contextmanager
def db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def init_db():
    DB_DIR.mkdir(exist_ok=True)
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS materials (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                source_text TEXT NOT NULL,
                file_name TEXT,
                language TEXT DEFAULT 'kk',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS adaptations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL UNIQUE,
                result_json TEXT NOT NULL,
                engine TEXT NOT NULL DEFAULT 'local',
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL,
                question TEXT NOT NULL,
                answer TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS progress (
                material_id INTEGER PRIMARY KEY,
                completed INTEGER NOT NULL DEFAULT 0,
                score INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('student','teacher')),
                subject TEXT,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS classes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                code TEXT NOT NULL UNIQUE,
                created_at TEXT NOT NULL,
                FOREIGN KEY(teacher_id) REFERENCES users(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS class_members (
                class_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                joined_at TEXT NOT NULL,
                PRIMARY KEY(class_id, student_id),
                FOREIGN KEY(class_id) REFERENCES classes(id) ON DELETE CASCADE,
                FOREIGN KEY(student_id) REFERENCES users(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS quiz_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                material_id INTEGER NOT NULL,
                score INTEGER NOT NULL,
                total INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES users(id) ON DELETE CASCADE,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS missions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                class_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                description TEXT NOT NULL,
                due_date TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(teacher_id) REFERENCES users(id) ON DELETE CASCADE,
                FOREIGN KEY(class_id) REFERENCES classes(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS mission_completions (
                mission_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                completed_at TEXT NOT NULL,
                PRIMARY KEY(mission_id, student_id),
                FOREIGN KEY(mission_id) REFERENCES missions(id) ON DELETE CASCADE,
                FOREIGN KEY(student_id) REFERENCES users(id) ON DELETE CASCADE
            );
            """
        )


def now():
    return datetime.now().isoformat(timespec="seconds")


def clean_text(value):
    value = str(value or "").replace("\x00", " ").replace("\r", "")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def split_sentences(text):
    return [
        item.strip()
        for item in re.split(
            r"(?<=[.!?])\s+(?=[А-ЯӘҒҚҢӨҰҮҺІA-Z0-9])", clean_text(text)
        )
        if item.strip()
    ]


def get_words(text):
    return re.findall(
        r"[а-яәіңғқөұүһіa-z0-9'-]+", clean_text(text).lower(), flags=re.I
    )


def keywords(text, limit=8):
    counts = {}
    for word in get_words(text):
        if len(word) >= 4 and word not in STOP_WORDS:
            counts[word] = counts.get(word, 0) + 1
    return [
        word
        for word, _ in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[
            :limit
        ]
    ]


def select_summary(text, limit=4):
    sentence_list = split_sentences(text)
    if len(sentence_list) <= limit:
        return sentence_list
    terms = keywords(text, limit)
    scored = []
    for position, sentence in enumerate(sentence_list):
        score = sum(2 for term in terms if term in sentence.lower())
        score += 1 if position == 0 else 0
        scored.append((score, position, sentence))
    return [
        row[2]
        for row in sorted(sorted(scored, reverse=True)[:limit], key=lambda row: row[1])
    ]


def local_adaptation(text, language="kk"):
    source = clean_text(text)
    terms = keywords(source)
    summary_sentences = select_summary(source)
    topic = ", ".join(terms[:3]) or "негізгі тақырып"
    short = "\n".join(
        f"• {sentence[:180]}{'…' if len(sentence) > 180 else ''}"
        for sentence in summary_sentences
    )
    steps = "\n".join(
        f"{index}. {sentence}"
        for index, sentence in enumerate(split_sentences(source)[:8], 1)
    ) or "1. Мәтінді оқыңыз.\n2. Негізгі ұғымдарды белгілеңіз.\n3. Мысал арқылы тексеріңіз."
    examples = "\n".join(
        f"{index}. «{term}» ұғымын күнделікті өмірдегі жағдаймен байланыстырып, өз мысалыңызды жазыңыз."
        for index, term in enumerate(terms[:4], 1)
    ) or "1. Тақырыпқа қатысты қарапайым жағдай ойластырыңыз.\n2. Оны негізгі ұғыммен байланыстырыңыз."
    tasks = (
        f"1. «{terms[0] if terms else 'негізгі ұғым'}» ұғымын өз сөзіңізбен түсіндіріңіз.\n"
        f"2. {terms[1] if len(terms) > 1 else 'тақырып'} бойынша бір мысал келтіріңіз.\n"
        "3. Мәтіндегі екі негізгі ойдың байланысын жазыңыз.\n"
        "4. Жауабыңыздан негізгі терминдерді белгілеңіз."
    )
    plain = re.sub(r"\([^)]*\)", "", source).replace(";", ".")
    plain = "\n\n".join(
        sentence[:200] + ("…" if len(sentence) > 200 else "")
        for sentence in split_sentences(plain)
    )
    nodes = (terms or ["Негізгі ұғым", "Ереже", "Мысал", "Нәтиже"])[:5]
    return {
        "simplified": short or source[:600],
        "steps": steps,
        "examples": examples,
        "diagram": {
            "title": topic,
            "nodes": nodes,
            "edges": [[nodes[index], nodes[index + 1]] for index in range(len(nodes) - 1)],
        },
        "audio": f"{topic}. " + " ".join(summary_sentences),
        "large": plain or source,
        "summary": " ".join(summary_sentences) or source[:500],
        "tasks": tasks,
        "keywords": terms,
        "explanation": (
            f"Бұл материалдың негізгі бағыты: {topic}. Алдымен анықтаманы түсініңіз, "
            "кейін оның қолданылуын мысалмен тексеріңіз."
        ),
        "language": language,
    }


def ai_configured():
    return bool(os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY"))


def ask_model(system, prompt):
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY")
    if not api_key:
        return None
    endpoint = (
        os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
        + "/chat/completions"
    )
    payload = json.dumps(
        {
            "model": os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            "temperature": float(os.getenv("OPENAI_TEMPERATURE", "0.25")),
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
        }
    ).encode()
    model_request = urllib.request.Request(
        endpoint,
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    with urllib.request.urlopen(model_request, timeout=45) as response:
        data = json.loads(response.read().decode())
    return data["choices"][0]["message"]["content"]


def extract_file(file):
    extension = Path(file.filename or "").suffix.lower()
    allowed = {".txt", ".md", ".csv", ".html", ".pdf", ".docx", ".pptx"}
    if extension not in allowed:
        raise ValueError("PDF, DOCX, PPTX, TXT, MD, CSV немесе HTML файлдарын қолданыңыз.")
    data = file.read()
    if extension in {".txt", ".md", ".csv", ".html"}:
        return clean_text(data.decode("utf-8", errors="ignore"))
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return clean_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        return clean_text("\n\n".join(paragraph.text for paragraph in document.paragraphs))
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides = []
    for slide in presentation.slides:
        slide_text = [
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        ]
        if slide_text:
            slides.append(" ".join(slide_text))
    return clean_text("\n\n".join(slides))


def local_answer(question, source):
    query = keywords(question, 8)
    matched = []

    for position, sentence in enumerate(split_sentences(source)):
        score = sum(1 for term in query if term in sentence.lower())

        if score:
            matched.append((score, -position, sentence))

    matched = [row[2] for row in sorted(matched, reverse=True)[:3]]

    if not matched:
        return "Бұл сұраққа жүктелген материалдан нақты жауап табылмады. Сұрақты мәтіндегі негізгі терминдермен нақтылап көріңіз."

    return "Материалға сүйенген жауап:\n\n" + "\n".join(
        f"• {sentence}" for sentence in matched
    )


def local_quiz(source, language="kk"):
    sentences = [
        sentence.strip()
        for sentence in split_sentences(source)
        if len(sentence.strip()) > 20
    ]

    if len(sentences) < 5:
        return [], "local"

    questions = []

    for i in range(5):
        correct_answer = sentences[i]

        options = [correct_answer]

        for j in range(len(sentences)):
            if j != i and sentences[j] not in options:
                options.append(sentences[j])

            if len(options) == 4:
                break

        questions.append({
            "question": f"Материал бойынша дұрыс жауапты таңдаңыз: {correct_answer[:100]}...",
            "options": options,
            "correct": 0
        })

    return questions, "local"


def create_quiz(source, language="kk"):
    if not source:
        return [], "local"

    if ai_configured():
        prompt = f"""
Create exactly 5 high-quality multiple-choice questions from the source material.

Rules:
- Use ONLY information explicitly supported by the source.
- Do not invent facts.
- Each question must test understanding, not just copy a sentence.
- Use exactly 4 options per question.
- Exactly one option must be correct.
- The correct field is a zero-based option index from 0 to 3.
- Write the questions and answers in {LANGUAGE_NAMES.get(language, 'қазақ тілі')}.
- Return ONLY valid JSON. No markdown and no extra text.

JSON format:
{{
  "quiz": [
    {{
      "question": "...",
      "options": ["...", "...", "...", "..."],
      "correct": 0
    }}
  ]
}}

SOURCE:
{source[:14000]}
"""

        try:
            answer = ask_model(
                "You are Ayla AI, a careful educational quiz generator. Ground every question in the provided source.",
                prompt
            )
            data = json.loads(answer)
            quiz = data.get("quiz") if isinstance(data, dict) else None
            if isinstance(quiz, list):
                clean_quiz = []
                for item in quiz[:5]:
                    if not isinstance(item, dict):
                        continue
                    question = clean_text(item.get("question"))
                    options = item.get("options")
                    correct = item.get("correct")
                    if question and isinstance(options, list) and len(options) == 4 and isinstance(correct, int) and 0 <= correct < 4:
                        clean_quiz.append({
                            "question": question,
                            "options": [clean_text(x) for x in options],
                            "correct": correct,
                        })
                if len(clean_quiz) == 5:
                    return clean_quiz, "ai"
        except Exception:
            pass

    return local_quiz(source, language)
    


def material_json(row, adaptation=None, progress=None, preview=False):
    result = dict(row)
    if preview:
        result["source_text"] = result["source_text"][:180]
    result["adaptation"] = json.loads(adaptation["result_json"]) if adaptation else None
    result["progress"] = dict(progress) if progress else {
        "completed": 0,
        "score": 0,
    }
    return result


@app.post("/api/auth/register")
def register():
    body = request.get_json(silent=True) or {}
    name = clean_text(body.get("name"))
    email = clean_text(body.get("email")).lower()
    password = str(body.get("password") or "")
    role = body.get("role") or "student"
    subject = clean_text(body.get("subject"))
    if not name or not email or len(password) < 6 or role not in {"student", "teacher"}:
        return jsonify({"error": "Аты, email, рөл және кемінде 6 таңбалы пароль қажет."}), 400
    try:
        with db() as connection:
            cur = connection.execute(
                "INSERT INTO users(name,email,password_hash,role,subject,created_at) VALUES(?,?,?,?,?,?)",
                (name, email, generate_password_hash(password), role, subject, now()),
            )
            user_id = cur.lastrowid
        session["user_id"] = user_id
        session["role"] = role
        return jsonify({"ok": True, "user": {"id": user_id, "name": name, "email": email, "role": role, "subject": subject}}), 201
    except sqlite3.IntegrityError:
        return jsonify({"error": "Бұл email бұрын тіркелген."}), 409


@app.post("/api/auth/login")
def login():
    body = request.get_json(silent=True) or {}
    email = clean_text(body.get("email")).lower()
    password = str(body.get("password") or "")
    with db() as connection:
        user = connection.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if not user or not check_password_hash(user["password_hash"], password):
        return jsonify({"error": "Email немесе пароль дұрыс емес."}), 401
    session["user_id"] = user["id"]
    session["role"] = user["role"]
    return jsonify({"ok": True, "user": {"id": user["id"], "name": user["name"], "email": user["email"], "role": user["role"], "subject": user["subject"]}})


@app.post("/api/auth/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/auth/me")
def me():
    user_id = session.get("user_id")
    if not user_id:
        return jsonify({"authenticated": False})
    with db() as connection:
        user = connection.execute("SELECT id,name,email,role,subject FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user:
        session.clear()
        return jsonify({"authenticated": False})
    return jsonify({"authenticated": True, "user": dict(user)})


def current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    with db() as connection:
        return connection.execute("SELECT id,name,email,role,subject FROM users WHERE id = ?", (user_id,)).fetchone()


def _student_streak(connection, student_id):
    rows = connection.execute(
        "SELECT DISTINCT substr(created_at,1,10) day FROM quiz_results WHERE student_id=? ORDER BY day DESC",
        (student_id,),
    ).fetchall()
    days = [datetime.fromisoformat(row["day"]).date() for row in rows if row["day"]]
    if not days:
        return 0
    streak = 1
    for index in range(1, len(days)):
        if (days[index - 1] - days[index]).days == 1:
            streak += 1
        else:
            break
    return streak


def _student_badges(quiz_count, average, streak):
    badges = []
    if quiz_count >= 1:
        badges.append({"icon": "✦", "title": "Алғашқы қадам", "text": "Бірінші Quiz аяқталды"})
    if quiz_count >= 5:
        badges.append({"icon": "⚡", "title": "Белсенді оқушы", "text": "5 Quiz аяқталды"})
    if average >= 80:
        badges.append({"icon": "◆", "title": "Жақсы нәтиже", "text": "Орташа нәтиже 80%+"})
    if streak >= 3:
        badges.append({"icon": "🔥", "title": "Үздіксіз оқу", "text": "3 күндік серия"})
    return badges


@app.get("/api/dashboard")
def dashboard():
    user = current_user()
    if not user:
        return jsonify({"error": "Кіру қажет."}), 401
    with db() as connection:
        if user["role"] == "student":
            results = connection.execute(
                "SELECT COUNT(*) total, COALESCE(AVG(score * 100.0 / NULLIF(total,0)),0) avg_score "
                "FROM quiz_results WHERE student_id = ?",
                (user["id"],),
            ).fetchone()
            classes = connection.execute(
                "SELECT c.id,c.name,c.code FROM classes c JOIN class_members m ON m.class_id=c.id "
                "WHERE m.student_id=? ORDER BY c.created_at DESC",
                (user["id"],),
            ).fetchall()
            recent = connection.execute(
                "SELECT qr.score,qr.total,qr.created_at,m.title FROM quiz_results qr "
                "JOIN materials m ON m.id=qr.material_id WHERE qr.student_id=? "
                "ORDER BY qr.created_at DESC LIMIT 5",
                (user["id"],),
            ).fetchall()
            quiz_count = int(results["total"] or 0)
            average = round(results["avg_score"] or 0)
            streak = _student_streak(connection, user["id"])
            badges = _student_badges(quiz_count, average, streak)
            today = now()[:10]
            teacher_mission = connection.execute(
                "SELECT m.id,m.title,m.description,m.due_date, "
                "EXISTS(SELECT 1 FROM mission_completions mc WHERE mc.mission_id=m.id AND mc.student_id=?) completed "
                "FROM missions m JOIN class_members cm ON cm.class_id=m.class_id "
                "WHERE cm.student_id=? AND m.due_date>=? ORDER BY m.created_at DESC LIMIT 1",
                (user["id"], user["id"], today),
            ).fetchone()
            if teacher_mission:
                mission = {
                    "source": "teacher",
                    "id": teacher_mission["id"],
                    "title": teacher_mission["title"],
                    "text": teacher_mission["description"],
                    "done": 1 if teacher_mission["completed"] else 0,
                    "target": 1,
                    "due_date": teacher_mission["due_date"],
                }
            else:
                mission_target = 1 if quiz_count == 0 else 2
                mission_done = min(quiz_count, mission_target)
                mission = {
                    "source": "system",
                    "id": None,
                    "title": "Ayla AI миссиясы",
                    "text": "1 Quiz орындаңыз" if quiz_count == 0 else "Тағы бір Quiz орындап, оқу серияңызды жалғастырыңыз",
                    "done": mission_done,
                    "target": mission_target,
                }
            return jsonify({
                "role": "student",
                "quiz_count": quiz_count,
                "average_score": average,
                "streak": streak,
                "badges": badges,
                "mission": mission,
                "classes": [dict(x) for x in classes],
                "recent": [dict(x) for x in recent],
            })

        classes = connection.execute(
            "SELECT c.id,c.name,c.code,(SELECT COUNT(*) FROM class_members m WHERE m.class_id=c.id) students "
            "FROM classes c WHERE c.teacher_id=? ORDER BY c.created_at DESC",
            (user["id"],),
        ).fetchall()
        total_students = 0
        class_data = []
        for item in classes:
            stats = connection.execute(
                "SELECT COUNT(qr.id) attempts, COALESCE(AVG(qr.score * 100.0 / NULLIF(qr.total,0)),0) avg_score "
                "FROM quiz_results qr JOIN class_members cm ON cm.student_id=qr.student_id "
                "WHERE cm.class_id=?",
                (item["id"],),
            ).fetchone()
            total_students += int(item["students"] or 0)
            class_data.append({**dict(item), "attempts": int(stats["attempts"] or 0), "average": round(stats["avg_score"] or 0)})
        missions = connection.execute(
            "SELECT m.id,m.class_id,m.title,m.description,m.due_date,c.name class_name "
            "FROM missions m JOIN classes c ON c.id=m.class_id WHERE m.teacher_id=? "
            "ORDER BY m.created_at DESC LIMIT 20", (user["id"],)
        ).fetchall()
        return jsonify({
            "role": "teacher",
            "classes": class_data,
            "class_count": len(class_data),
            "total_students": total_students,
            "missions": [dict(x) for x in missions],
        })


@app.post("/api/classes")
def create_class():
    user = current_user()
    if not user or user["role"] != "teacher":
        return jsonify({"error":"Мұғалім ретінде кіру қажет."}), 403
    body = request.get_json(silent=True) or {}
    name = clean_text(body.get("name"))
    code = re.sub(r"[^A-Z0-9]", "", str(body.get("code") or "").upper()) or os.urandom(3).hex().upper()
    if not name:
        return jsonify({"error":"Сынып атауын енгізіңіз."}), 400
    try:
        with db() as connection:
            cur=connection.execute("INSERT INTO classes(teacher_id,name,code,created_at) VALUES(?,?,?,?)",(user["id"],name,code,now()))
            cid=cur.lastrowid
        return jsonify({"id":cid,"name":name,"code":code}),201
    except sqlite3.IntegrityError:
        return jsonify({"error":"Бұл код қолданылып қойған."}),409


@app.post("/api/classes/join")
def join_class():
    user=current_user()
    if not user or user["role"] != "student":
        return jsonify({"error":"Оқушы ретінде кіру қажет."}),403
    body=request.get_json(silent=True) or {}
    code=clean_text(body.get("code")).upper()
    with db() as connection:
        cls=connection.execute("SELECT * FROM classes WHERE code=?",(code,)).fetchone()
        if not cls:
            return jsonify({"error":"Сынып коды табылмады."}),404
        connection.execute("INSERT OR IGNORE INTO class_members(class_id,student_id,joined_at) VALUES(?,?,?)",(cls["id"],user["id"],now()))
    return jsonify({"ok":True,"class":{"id":cls["id"],"name":cls["name"],"code":cls["code"]}})


@app.get("/api/classes/<int:class_id>/students")
def class_students(class_id):
    user=current_user()
    if not user or user["role"] != "teacher":
        return jsonify({"error":"Рұқсат жоқ."}),403
    with db() as connection:
        owner=connection.execute("SELECT id FROM classes WHERE id=? AND teacher_id=?",(class_id,user["id"])).fetchone()
        if not owner: return jsonify({"error":"Сынып табылмады."}),404
        rows=connection.execute("SELECT u.id,u.name,u.email FROM users u JOIN class_members m ON m.student_id=u.id WHERE m.class_id=? ORDER BY u.name",(class_id,)).fetchall()
    return jsonify({"students":[dict(x) for x in rows]})


@app.post("/api/quiz-results")
def save_quiz_result():
    user=current_user()
    if not user or user["role"] != "student": return jsonify({"error":"Оқушы ретінде кіру қажет."}),403
    body=request.get_json(silent=True) or {}
    material_id=int(body.get("material_id",0)); score=int(body.get("score",0)); total=max(1,int(body.get("total",1)))
    with db() as connection:
        exists=connection.execute("SELECT id FROM materials WHERE id=?",(material_id,)).fetchone()
        if not exists: return jsonify({"error":"Материал табылмады."}),404
        connection.execute("INSERT INTO quiz_results(student_id,material_id,score,total,created_at) VALUES(?,?,?,?,?)",(user["id"],material_id,score,total,now()))
    return jsonify({"ok":True,"score":score,"total":total})


@app.post("/api/missions")
def create_mission():
    user = current_user()
    if not user or user["role"] != "teacher":
        return jsonify({"error": "Мұғалім ретінде кіру қажет."}), 403
    body = request.get_json(silent=True) or {}
    title = clean_text(body.get("title"))
    description = clean_text(body.get("description"))
    class_id = int(body.get("class_id", 0))
    due_date = clean_text(body.get("due_date")) or now()[:10]
    if not title or not description or not class_id:
        return jsonify({"error": "Миссия атауын, сипаттамасын және сыныпты енгізіңіз."}), 400
    with db() as connection:
        owned = connection.execute("SELECT id FROM classes WHERE id=? AND teacher_id=?", (class_id, user["id"])).fetchone()
        if not owned:
            return jsonify({"error": "Сынып табылмады."}), 404
        cur = connection.execute(
            "INSERT INTO missions(teacher_id,class_id,title,description,due_date,created_at) VALUES(?,?,?,?,?,?)",
            (user["id"], class_id, title, description, due_date, now()),
        )
        mission_id = cur.lastrowid
    return jsonify({"id": mission_id, "title": title, "description": description, "due_date": due_date}), 201


@app.post("/api/missions/<int:mission_id>/complete")
def complete_mission(mission_id):
    user = current_user()
    if not user or user["role"] != "student":
        return jsonify({"error": "Оқушы ретінде кіру қажет."}), 403
    with db() as connection:
        mission = connection.execute(
            "SELECT m.id FROM missions m JOIN class_members cm ON cm.class_id=m.class_id "
            "WHERE m.id=? AND cm.student_id=?", (mission_id, user["id"])
        ).fetchone()
        if not mission:
            return jsonify({"error": "Миссия табылмады."}), 404
        connection.execute(
            "INSERT OR REPLACE INTO mission_completions(mission_id,student_id,completed_at) VALUES(?,?,?)",
            (mission_id, user["id"], now()),
        )
    return jsonify({"ok": True})


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "database": DB_PATH.exists(), "ai_configured": ai_configured()})


@app.get("/api/materials")
def list_materials():
    query = clean_text(request.args.get("q", ""))
    with db() as connection:
        if query:
            rows = connection.execute(
                "SELECT * FROM materials WHERE title LIKE ? OR source_text LIKE ? "
                "ORDER BY updated_at DESC LIMIT 100",
                (f"%{query}%", f"%{query}%"),
            ).fetchall()
        else:
            rows = connection.execute(
                "SELECT * FROM materials ORDER BY updated_at DESC LIMIT 100"
            ).fetchall()
        materials = []
        for row in rows:
            progress = connection.execute(
                "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
                (row["id"],),
            ).fetchone()
            materials.append(material_json(row, progress=progress, preview=True))
    return jsonify({"materials": materials})


@app.get("/api/materials/<int:material_id>")
def get_material(material_id):
    with db() as connection:
        row = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        adaptation = connection.execute(
            "SELECT * FROM adaptations WHERE material_id = ?", (material_id,)
        ).fetchone()
        progress = connection.execute(
            "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
            (material_id,),
        ).fetchone()
    if not row:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify(material_json(row, adaptation, progress))


@app.post("/api/materials/<int:material_id>/quiz")
def generate_material_quiz(material_id):

    body = request.get_json(
        silent=True
    ) or {}

    requested_language = (
        body.get("language")
        or "kk"
    )

    with db() as connection:

        material = connection.execute(
            "SELECT * FROM materials WHERE id = ?",
            (material_id,)
        ).fetchone()

    if not material:

        return jsonify({
            "error": "Материал табылмады."
        }), 404

    if requested_language not in LANGUAGE_NAMES:
        requested_language = (
            material["language"]
            or "kk"
        )

    quiz, engine = create_quiz(
        material["source_text"],
        requested_language
    )

    if not quiz:

        return jsonify({
            "error":
                "Материал бойынша тест "
                "жасауға мәтін жеткіліксіз."
        }), 400

    return jsonify({
        "quiz": quiz,
        "engine": engine,
        "language": requested_language
    })


@app.post("/api/materials")
def create_material():
    try:
        body = request.get_json(silent=True) if request.is_json else {}
        title = clean_text(request.form.get("title") or body.get("title", ""))
        language = request.form.get("language") or body.get("language") or "kk"
        file = request.files.get("file")
        text = request.form.get("text", "") or body.get("text", "")
        file_name = file.filename if file else ""
        if file:
            text = extract_file(file)
        text = clean_text(text)
        if not text:
            return jsonify({"error": "Материал мәтінін енгізіңіз немесе файл жүктеңіз."}), 400
        if not title:
            title = Path(file_name).stem if file_name else "Жаңа оқу материалы"
        timestamp = now()
        with db() as connection:
            cursor = connection.execute(
                "INSERT INTO materials(title, source_text, file_name, language, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (title, text, file_name, language, timestamp, timestamp),
            )
            material_id = cursor.lastrowid
        result = await_adaptation(material_id, text, language)
        return jsonify(
            {"id": material_id, "title": title, "source_text": text, "adaptation": result}
        ), 201
    except Exception as error:
        return jsonify({"error": str(error)}), 400


def await_adaptation(material_id, text, language):
    result = local_adaptation(text, language)
    engine = "local"
    if ai_configured():
        try:
            prompt = f"""
Сен IncluLearn AI платформасындағы Ayla AI білім беру ассистентісің.
Берілген оқу материалын оқушыға түсінікті және инклюзивті форматқа бейімде.

Қатаң ережелер:
1. Тек берілген материалдағы ақпаратқа сүйен. Ойдан дерек, факт, анықтама қоспа.
2. Маңызды терминдерді сақта, бірақ күрделі сөйлемдерді қарапайым тілмен түсіндір.
3. Жауап {LANGUAGE_NAMES.get(language, 'қазақ тілі')} тілінде болсын.
4. Оқушыға арналған нақты, қысқа және пайдалы мәтін жаса.
5. steps нөмірленген қадамдар болсын.
6. examples материалдағы ұғымдарды түсінуге көмектесетін мысалдар болсын; материалда жоқ фактіні шындық ретінде қоспа.
7. tasks оқушы орындай алатын нақты тапсырмалар болсын.
8. audio табиғи түрде тыңдауға болатын тұтас мәтін болсын.
9. summary тек негізгі ойларды қамтысын.

Тек мына JSON құрылымын қайтар:
{{
  "simplified": "...",
  "steps": "1. ...\n2. ...",
  "examples": "1. ...\n2. ...",
  "explanation": "...",
  "summary": "...",
  "tasks": "1. ...\n2. ...",
  "audio": "..."
}}

МАТЕРИАЛ:
{text[:16000]}
"""
            response = ask_model(
                "Сен IncluLearn AI платформасындағы Ayla AI ассистентісің. Қысқа, түсінікті жауап бер.",
                prompt,
            )
            parsed = json.loads(response)
            result.update(
                {key: parsed[key] for key in parsed if key in result and isinstance(parsed[key], str)}
            )
            engine = "ai"
        except Exception:
            engine = "local"
    with db() as connection:
        connection.execute(
            "INSERT INTO adaptations(material_id, result_json, engine, created_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET result_json=excluded.result_json, "
            "engine=excluded.engine, created_at=excluded.created_at",
            (material_id, json.dumps(result, ensure_ascii=False), engine, now()),
        )
        connection.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?", (now(), material_id)
        )
    return result


@app.put("/api/materials/<int:material_id>")
def update_material(material_id):
    body = request.get_json(silent=True) or {}
    title = clean_text(body.get("title"))
    text = clean_text(body.get("text"))
    if not title or not text:
        return jsonify({"error": "Атауы мен мәтіні бос болмауы керек."}), 400
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "UPDATE materials SET title = ?, source_text = ?, updated_at = ? WHERE id = ?",
            (title, text, now(), material_id),
        )
    result = await_adaptation(material_id, text, body.get("language", "kk"))
    return jsonify({"id": material_id, "title": title, "source_text": text, "adaptation": result})


@app.delete("/api/materials/<int:material_id>")
def delete_material(material_id):
    with db() as connection:
        cursor = connection.execute("DELETE FROM materials WHERE id = ?", (material_id,))
    if cursor.rowcount == 0:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify({"ok": True})


@app.post("/api/materials/<int:material_id>/progress")
def save_progress(material_id):
    body = request.get_json(silent=True) or {}
    completed = 1 if body.get("completed") else 0
    score = max(0, min(100, int(body.get("score", 0))))
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "INSERT INTO progress(material_id, completed, score, updated_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET completed=excluded.completed, "
            "score=excluded.score, updated_at=excluded.updated_at",
            (material_id, completed, score, now()),
        )
    return jsonify({"completed": completed, "score": score})


@app.post("/api/materials/<int:material_id>/ask")
def ask_about_material(material_id):
    body = request.get_json(silent=True) or {}
    question = clean_text(body.get("question"))
    if not question:
        return jsonify({"error": "Сұрақ жазыңыз."}), 400
    with db() as connection:
        material = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
    if not material:
        return jsonify({"error": "Материал табылмады."}), 404
    answer = local_answer(question, material["source_text"])
    engine = "local"
    if ai_configured():
        try:
            answer = ask_model(
                """Сен Ayla AI — IncluLearn AI платформасындағы жеке оқу ассистентісің.
Жауапты оқушыға түсінікті, нақты және пайдалы етіп бер.
Қатаң ережелер:
1. Негізгі жауапты тек берілген оқу материалына сүйеніп құрастыр.
2. Материалда жоқ фактіні ойдан қоспа. Егер жауап материалда жоқ болса, оны ашық айт.
3. Алдымен қысқа тікелей жауап бер, кейін қажет болса 2–4 қысқа түсіндіру немесе мысал келтір.
4. Күрделі термин болса, оны қарапайым тілмен түсіндір.
5. Оқушыға жасына сай, жылы, бірақ академиялық стильде жауап бер.
6. Сұрақ тапсырма сұраса, қадамдармен көрсет.
7. Сұрақ түсініксіз болса, нақтылау сұрағын қой.
8. Жауапты тек {language} тілінде бер.
""".format(language=LANGUAGE_NAMES.get(material['language'], 'қазақ тілі')),
                f"ОҚУ МАТЕРИАЛЫ:\n{material['source_text']}\n\nОҚУШЫНЫҢ СҰРАҒЫ:\n{question}",
            )
            engine = "ai"
        except Exception:
            pass
    with db() as connection:
        connection.execute(
            "INSERT INTO questions(material_id, question, answer, created_at) VALUES (?, ?, ?, ?)",
            (material_id, question, answer, now()),
        )
    return jsonify({"answer": answer, "engine": engine})


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "3000")), debug=False)